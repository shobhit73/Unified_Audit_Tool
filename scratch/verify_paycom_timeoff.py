"""Verify the Paycom Time Off tool against the real Banda Logistics files.

Run:  python scratch/verify_paycom_timeoff.py [banda|match|before|synthetic|ui ...]

Spec: docs/superpowers/specs/2026-10-06-paycom-timeoff-parity-design.md

The old tool is loaded from a PINNED commit (`ace3c1d`, main before this
change), never from HEAD, so the before/after comparison stays meaningful once
this branch merges.
"""
import importlib.util
import io
import json
import os
import subprocess
import sys
import tempfile

import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from apps.paycom import timeoff_audit as ta          # noqa: E402

BASELINE = "ace3c1d"
B = r"C:\Users\rohit.kaushik\Downloads\Banda Logistics"
PAYCOM = B + r"\20261002123220_TimeOff_Summary_Report_3fe976635c3ba3363caa4862ac801a6f6e2ebe7c.xlsx"
TEMPLATE = B + r"\Time Off Import.xlsx"
CENSUS = B + r"\Multi_Client_Banda Logistics LLC_Employee_Census.xlsm"

failures = 0


class U(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name
        self.size = len(data)


def up(path_or_bytes, name=None):
    if isinstance(path_or_bytes, bytes):
        return U(path_or_bytes, name or "upload.xlsx")
    return U(open(path_or_bytes, "rb").read(), os.path.basename(path_or_bytes))


def check(label, ok, detail=""):
    global failures
    print("   %-72s %s%s" % (label, "OK" if ok else "FAIL",
                             "" if ok else "  <- " + str(detail)[:300]))
    failures += 0 if ok else 1


def baseline():
    src = subprocess.run(["git", "-C", ROOT, "show", BASELINE + ":apps/paycom/timeoff_audit.py"],
                         capture_output=True, text=True, encoding="utf-8", check=True).stdout
    path = os.path.join(tempfile.gettempdir(), "paycom_timeoff_baseline.py")
    open(path, "w", encoding="utf-8").write(src)
    spec = importlib.util.spec_from_file_location("paycom_timeoff_baseline", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def paycom_df():
    return pd.read_excel(PAYCOM, dtype=str)


def paycom_bytes(df):
    buf = io.BytesIO()
    df.to_excel(buf, index=False, sheet_name="Time-Off Summary ")
    return buf.getvalue()


def template_bytes(edit=None):
    wb = openpyxl.load_workbook(TEMPLATE)
    if edit:
        edit(ta._template_sheet(wb))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def run(p=PAYCOM, t=TEMPLATE, mapping=None, include_salaried=False, include_blank_hourly=True):
    """Read + plan + fill + audit without Streamlit."""
    pdf, err = ta.read_paycom_balances(up(p, "paycom.xlsx"))
    assert not err, err
    tpl, err = ta.read_template_rows(up(t, "template.xlsx"))
    assert not err, err
    census, err = ta.read_census(up(CENSUS))
    assert not err, err
    auto = ta.auto_map([x for x, _, _ in ta.policy_summary(pdf)], tpl["policies"])
    mapping = mapping or auto
    plan = ta.plan_fill(tpl, pdf, mapping, include_salaried, include_blank_hourly)
    wb, filled, err = ta.fill_import_template(up(t, "template.xlsx"), plan["decisions"])
    assert not err, err
    sheets, counts = ta.build_audit_sheets(up(t, "template.xlsx"), tpl, pdf, census,
                                           plan, mapping, auto)
    return {"paycom": pdf, "tpl": tpl, "auto": auto, "plan": plan, "wb": wb,
            "filled": filled, "sheets": sheets, "counts": counts}


def openings(wb):
    """{Employee ID: Opening Balance} from the template sheet."""
    ws = ta._template_sheet(wb)
    head = [c.value for c in ws[ta.UZIO_HEADER_ROW]]
    i_id = head.index("Employee ID") + 1
    i_ob = head.index("Opening Balance") + 1
    return {str(ws.cell(row=r, column=i_id).value).strip(): ws.cell(row=r, column=i_ob).value
            for r in range(ta.UZIO_HEADER_ROW + 1, ws.max_row + 1)
            if ws.cell(row=r, column=i_id).value not in (None, "")}


def fnum(v):
    return float(str(v).replace(",", ""))


# ---------------------------------------------------------------- banda
def group_banda():
    print("\n== Banda, default settings (Salaried off, blank Hourly on)")
    p = paycom_df().set_index("EECode")
    r = run()
    tpl_rows = r["tpl"]["rows"]
    out = openings(r["wb"])
    before = openings(openpyxl.load_workbook(TEMPLATE))

    hourly_matched = [t for t in tpl_rows if t["raw_id"] in p.index and not t["salaried"]]
    salaried_matched = [t for t in tpl_rows if t["raw_id"] in p.index and t["salaried"]]
    unmatched = [t for t in tpl_rows if t["raw_id"] not in p.index]

    check("auto-map: Paid Time Off -> Paid PTO", r["auto"] == {"Paid Time Off": "Paid PTO"},
          r["auto"])
    check("rows written = matched Hourly rows (%d)" % len(hourly_matched),
          r["filled"] == len(hourly_matched), r["filled"])
    bad = [(t["raw_id"], out[t["raw_id"]], p.loc[t["raw_id"], "Available"])
           for t in hourly_matched
           if abs(fnum(out[t["raw_id"]]) - fnum(p.loc[t["raw_id"], "Available"])) > 0.001]
    check("every written cell = Paycom Available", not bad, bad[:5])
    differs = [t for t in hourly_matched
               if abs(fnum(p.loc[t["raw_id"], "Available"])
                      - fnum(p.loc[t["raw_id"], "Net Available"])) > 0.001]
    check("Net Available never written (%d rows where it differs)" % len(differs),
          all(abs(fnum(out[t["raw_id"]]) - fnum(p.loc[t["raw_id"], "Net Available"])) > 0.001
              for t in differs))
    check("%d Salaried rows untouched" % len(salaried_matched),
          all(out[t["raw_id"]] == before[t["raw_id"]] for t in salaried_matched))
    check("%d template-only IDs untouched" % len(unmatched),
          all(out[t["raw_id"]] == before[t["raw_id"]] for t in unmatched))

    exc = r["sheets"]["Exception Summary"]
    sal = exc[exc["Issue Category"] == "Salaried — not filled (Paid PTO)"]
    check("each skipped Salaried row is in Exception Summary",
          sorted(sal["Employee ID"]) == sorted(t["raw_id"] for t in salaried_matched),
          list(sal["Employee ID"]))
    check("Exception Summary has the Paycom Balance column",
          list(exc.columns) == ta.EXCEPTION_COLUMNS, list(exc.columns))
    check("no future-time-off category in Exception Summary",
          not exc["Issue Category"].str.contains("uture").any())

    fut = r["sheets"]["Future Time Off"]
    raw = paycom_df()
    n_future = int(((pd.to_numeric(raw["Future Approved"]) > 0)
                    | (pd.to_numeric(raw["Future Pending"]) > 0)).sum())
    check("Future Time Off sheet: %d rows" % n_future, len(fut) == n_future, len(fut))
    check("Future Time Off columns", list(fut.columns) == ta.FUTURE_COLUMNS, list(fut.columns))
    abel = fut[fut["Employee ID"] == "A03K"]
    check("ABEL: Available 4.98, Future Approved 53.53, Net -48.55",
          len(abel) == 1 and abel.iloc[0]["Available"] == 4.98
          and abel.iloc[0]["Future Approved"] == 53.53
          and abel.iloc[0]["Net Available"] == -48.55, abel.to_dict("records"))
    check("future count = distinct employees", r["counts"]["future"] == fut["Employee ID"].nunique())

    pm = r["sheets"]["Policy Mapping"]
    check("Policy Mapping: one row, Auto",
          pm.to_dict("records")[0]["Set By"] == "Auto" and len(pm) == 1, pm.to_dict("records"))
    stat = r["sheets"]["Balance vs UZIO Status"]
    rank = stat["UZIO Employment Status"].str.lower().map(
        lambda s: 0 if s.startswith("terminated") else (2 if s == "active" else 1))
    check("Balance vs UZIO Status: Terminated first", rank.is_monotonic_increasing)
    check("census status reaches Exception Summary",
          (exc["Employment Status"] != "(not in census)").any())

    # the import file is the template, bar Opening Balance cells
    src = openpyxl.load_workbook(TEMPLATE)
    check("import file keeps only the template's own sheets",
          r["wb"].sheetnames == src.sheetnames, r["wb"].sheetnames)
    ob_col = [c.value for c in ta._template_sheet(src)[ta.UZIO_HEADER_ROW]].index("Opening Balance") + 1
    diffs = []
    for s_src, s_out in zip(src.worksheets, r["wb"].worksheets):
        for row in s_src.iter_rows():
            for c in row:
                v = s_out.cell(row=c.row, column=c.column).value
                if v != c.value and not (s_src.title == ta._template_sheet(src).title
                                         and c.column == ob_col and c.row > ta.UZIO_HEADER_ROW):
                    diffs.append((s_src.title, c.coordinate, c.value, v))
    check("no cell changed outside Opening Balance", not diffs, diffs[:5])

    print("\n== Banda, Salaried box ticked")
    r2 = run(include_salaried=True)
    out2 = openings(r2["wb"])
    check("Salaried rows now = Available",
          all(abs(fnum(out2[t["raw_id"]]) - fnum(p.loc[t["raw_id"], "Available"])) < 0.001
              for t in salaried_matched))
    check("rows written = all matched rows (%d)" % (len(hourly_matched) + len(salaried_matched)),
          r2["filled"] == len(hourly_matched) + len(salaried_matched), r2["filled"])


# ---------------------------------------------------------------- possible match
def group_match():
    print("\n== Possible Census Match (ID not in census -> look up by name)")
    r = run()
    exc = r["sheets"]["Exception Summary"]
    stat = r["sheets"]["Balance vs UZIO Status"]
    pm = ta.POSSIBLE_MATCH
    # Banda: three rehires sit in the census under a different EECode
    expected = {"A09S": "A09A (TERMINATED)", "A0H0": "A0I4 (TERMINATED)",
                "A0HW": "A0I5 (TERMINATED)"}
    no_match = ["A00S", "A03L", "A03M", "A049", "A04F", "A0BS", "A0JY", "A0K2"]
    for eid, want in expected.items():
        got_s = stat.loc[stat["Employee ID"] == eid, pm].tolist()
        got_e = exc.loc[exc["Employee ID"] == eid, pm].tolist()
        check("%s -> %s (status sheet + exceptions)" % (eid, want),
              got_s == [want] and got_e and set(got_e) == {want}, (got_s, got_e))
    got = stat.loc[stat["Employee ID"].isin(no_match), pm].tolist()
    check("the other 8 not-in-census employees stay blank (no surname-only match)",
          len(got) == 8 and all(g == "" for g in got), got)
    in_census = stat[stat["UZIO Employment Status"] != "(not in census)"]
    check("employees found by ID never get a suggestion", (in_census[pm] == "").all())
    check("column present in all three sheets",
          pm in exc.columns and pm in stat.columns
          and pm in r["sheets"]["Future Time Off"].columns)

    print("\n== name keys")
    check("'DROESE, ALEXANDRIA' == census Alexandria / Droese",
          ta._name_key(*ta._split_name("DROESE, ALEXANDRIA")) == ta._name_key("Alexandria", "Droese"))
    check("middle initial ignored", ta._name_key(*ta._split_name("SMITH, JOHN Q"))
          == ta._name_key("John", "Smith"))
    check("template 'First Last' form", ta._name_key(*ta._split_name("John Smith"))
          == ta._name_key("JOHN", "SMITH"))
    check("multi-word surname", ta._name_key(*ta._split_name("DE LA CRUZ, JUAN"))
          == ta._name_key(*ta._split_name("Juan De La Cruz")))
    check("a lone word is no key", ta._name_key(*ta._split_name("Cher")) is None)


# ---------------------------------------------------------------- before / after
def group_before():
    print("\n== before (%s) vs after, same Banda files" % BASELINE)
    old = baseline()
    old_bytes = old.run_tool(up(PAYCOM), up(TEMPLATE))
    check("old tool ran", bool(old_bytes))
    old_wb = openpyxl.load_workbook(io.BytesIO(old_bytes))
    o, n = openings(old_wb), openings(run()["wb"])
    p = paycom_df().set_index("EECode")
    rows = []
    for eid in n:
        if eid in p.index and o[eid] != n[eid]:
            rows.append((eid, p.loc[eid, "Employee"], fnum(p.loc[eid, "Available"]),
                         fnum(p.loc[eid, "Net Available"]), o[eid], n[eid]))
    tpl = {t["raw_id"]: t for t in run()["tpl"]["rows"]}
    old_net = all(abs(fnum(o[e]) - fnum(p.loc[e, "Net Available"])) < 0.001
                  for e in n if e in p.index)
    check("old tool wrote Net Available into every matched row", old_net)
    check("old tool also wrote the Salaried rows",
          all(o[e] != 0 or fnum(p.loc[e, "Net Available"]) == 0
              for e, t in tpl.items() if t["salaried"] and e in p.index))
    check("old import file carried 5 audit tabs",
          len(old_wb.sheetnames) == 7, old_wb.sheetnames)
    print("\n   %d rows differ. First 10:" % len(rows))
    print("   %-9s %-24s %10s %10s %10s %10s" % ("EECode", "Name", "Available", "Net Avail",
                                                "OLD", "NEW"))
    for row in rows[:10]:
        print("   %-9s %-24s %10.2f %10.2f %10s %10s" % (row[0], row[1][:24], row[2], row[3],
                                                         row[4], row[5]))
    sal_diff = sum(1 for x in rows if tpl[x[0]]["salaried"])
    print("   (of which %d Salaried rows: old wrote Net Available, new leaves them)" % sal_diff)


# ---------------------------------------------------------------- synthetic
def group_synthetic():
    print("\n== a second Time-Off Type starts on Do not import")
    raw = paycom_df()
    base = run()
    hourly = [t["raw_id"] for t in base["tpl"]["rows"]
              if not t["salaried"] and t["raw_id"] in set(raw["EECode"])][:5]
    extra = raw[raw["EECode"].isin(hourly)].copy()
    extra["Time-Off Type"] = "Sick"
    extra["Available"] = ["7.50", "7.50", "7.50", "7.50", "0.00"]
    extra["Future Approved"] = "0.00"
    extra["Future Pending"] = "0.00"
    two = paycom_bytes(pd.concat([raw, extra], ignore_index=True))
    r = run(p=two)
    check("Sick auto-maps to Do not import",
          r["auto"] == {"Paid Time Off": "Paid PTO", "Sick": ta.DO_NOT_IMPORT}, r["auto"])
    check("filled template identical to the one-type run",
          openings(r["wb"]) == openings(base["wb"]))
    exc = r["sheets"]["Exception Summary"]
    sick = exc[exc["Issue Category"] == "Paycom policy not imported (Sick)"]
    check("4 non-zero Sick balances reported, the 0.00 one not", len(sick) == 4, len(sick))

    print("\n== a blank Hourly Opening Balance")
    target = hourly[0]
    avail = fnum(raw.set_index("EECode").loc[target, "Available"])

    def blank_one(ws):
        head = [c.value for c in ws[ta.UZIO_HEADER_ROW]]
        for rr in range(ta.UZIO_HEADER_ROW + 1, ws.max_row + 1):
            if ws.cell(row=rr, column=head.index("Employee ID") + 1).value == target:
                ws.cell(row=rr, column=head.index("Opening Balance") + 1).value = None

    tb = template_bytes(blank_one)
    on = run(t=tb)
    check("filled by default with Available", openings(on["wb"])[target] == avail,
          openings(on["wb"])[target])
    exc = on["sheets"]["Exception Summary"]
    check("named as 'Blank filled — policy assigned (Paid PTO)'",
          ((exc["Employee ID"] == target)
           & (exc["Issue Category"] == "Blank filled — policy assigned (Paid PTO)")).any())
    off = run(t=tb, include_blank_hourly=False)
    check("left blank with the box off", openings(off["wb"])[target] is None,
          openings(off["wb"])[target])
    check("and listed under Unassigned Policies",
          target in set(off["sheets"]["Unassigned Policies"]["Employee ID"].astype(str)))

    print("\n== robustness")
    junk = raw.copy()
    junk.loc[0, "Future Approved"] = " "
    junk.loc[1, "Unit of Time"] = "Days"
    pdf, err = ta.read_paycom_balances(up(paycom_bytes(junk), "p.xlsx"))
    check("blank Future Approved does not crash", err is None and pdf is not None, err)
    check("blank Future Approved counts as 0",
          pdf is not None and pdf.set_index("id").loc[ta.clean_id(junk.loc[0, "EECode"]),
                                                      "future_approved"] == 0)
    check("Unit of Time = Days is reported",
          pdf is not None and ta.non_hour_units(pdf) == [("Paid Time Off", "Days", 1)],
          pdf is not None and ta.non_hour_units(pdf))
    pdf, err = ta.read_paycom_balances(up(paycom_bytes(raw.drop(columns=["Available"])), "p.xlsx"))
    check("a report with only Net Available is refused",
          pdf is None and "Available" in err and "Net Available" in err, err)
    csv = up(raw.to_csv(index=False).encode("utf-8"), "paycom.csv")
    pdf, err = ta.read_paycom_balances(csv)
    check("the same report as CSV reads the same",
          err is None and pdf["balance"].tolist() == run()["paycom"]["balance"].tolist(), err)


# ---------------------------------------------------------------- ui
def group_ui():
    print("\n== the screen")
    from streamlit.testing.v1 import AppTest

    def script():
        import io as _io
        import json as _json
        import os as _os
        import sys as _sys
        _sys.path.insert(0, _os.environ["PTO_REPO_ROOT"])
        import streamlit as _st
        from apps.paycom import timeoff_audit
        files = _json.loads(_os.environ.get("PTO_FILES", "{}"))

        class _U(_io.BytesIO):
            def __init__(self, p):
                with open(p, "rb") as fh:
                    super().__init__(fh.read())
                self.name = _os.path.basename(p)
                self.size = len(self.getvalue())

        _st.file_uploader = lambda *a, **k: (_U(files[k["key"]]) if k.get("key") in files
                                             else None)
        timeoff_audit.render_ui()

    os.environ["PTO_REPO_ROOT"] = ROOT

    os.environ["PTO_FILES"] = json.dumps({"pt_p": PAYCOM, "pt_u": TEMPLATE})
    at = AppTest.from_function(script, default_timeout=180).run()
    at.button(key="run_timeoff_paycom").click().run()
    check("Generate without the census names it",
          any("Please upload: UZIO Census." in e.value for e in at.error),
          [e.value for e in at.error])
    check("and offers no downloads", not at.get("download_button"))

    os.environ["PTO_FILES"] = json.dumps({"pt_p": PAYCOM, "pt_u": TEMPLATE, "pt_c": CENSUS})
    at = AppTest.from_function(script, default_timeout=180).run()
    check("one mapping dropdown, defaulting to Paid PTO",
          len(at.selectbox) == 1 and at.selectbox[0].value == "Paid PTO",
          [(s.label, s.value) for s in at.selectbox])
    sal = [cb for cb in at.checkbox if "salaried" in cb.label.lower()]
    blank = [cb for cb in at.checkbox if "blank" in cb.label.lower()]
    check("Salaried box unchecked, blank-Hourly box checked",
          len(sal) == 1 and sal[0].value is False and len(blank) == 1 and blank[0].value is True,
          [(cb.label, cb.value) for cb in at.checkbox])
    at.button(key="run_timeoff_paycom").click().run()
    check("green message, no error", len(at.success) == 1 and not at.error,
          [e.value for e in at.error])
    check("two download buttons", len(at.get("download_button")) == 2)
    check("the future-time-off line is shown",
          any("future time off" in i.value for i in at.info), [i.value for i in at.info])
    at.selectbox[0].set_value(ta.DO_NOT_IMPORT).run()
    check("changing the mapping discards the downloads", not at.get("download_button"))


GROUPS = {"banda": group_banda, "match": group_match, "before": group_before,
          "synthetic": group_synthetic, "ui": group_ui}


def main(argv):
    for name in (argv or list(GROUPS)):
        try:
            GROUPS[name]()
        except Exception as e:
            import traceback
            traceback.print_exc()
            check("group %s ran" % name, False, repr(e))
    print()
    print("FAIL - %d check(s)" % failures if failures else "PASS - every check held")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
