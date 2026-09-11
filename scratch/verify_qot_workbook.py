"""Task 4 — writing the filled template."""
import io
import os
import sys

import openpyxl
import pandas as pd

sys.path.insert(0, r"C:\Users\rohit.kaushik\Downloads\Unified_Audit_Tool")
import apps.adp.qualified_overtime as Q

CDC = r"C:\Users\rohit.kaushik\Downloads\CDC"
TEMPLATE = CDC + r"\Qualified_Overtime_Template_2026.xlsx"
PAYROLL = [CDC + r"\Payroll\Cleaned\Payroll History_Q3_cleaned.csv",
           CDC + r"\Payroll\Cleaned\PriorPayroll_03222026_06202026_06262026_cleaned.csv",
           CDC + r"\Payroll\Cleaned\PriorPayroll_12212025_03212026_03272026_cleaned.csv"]


class Upload(io.BytesIO):
    def __init__(self, path):
        super().__init__(open(path, "rb").read())
        self.name = os.path.basename(path)
        self.size = self.getbuffer().nbytes


ok = True


def check(label, got, want):
    global ok
    good = got == want
    ok = ok and good
    print("   %-50s %-46r %s" % (label, got, "OK" if good else "FAIL, want %r" % (want,)))


pay, _ = Q.read_payroll_files([Upload(p) for p in PAYROLL])
tpl, _ = Q.read_qot_template(Upload(TEMPLATE))
col, _ = Q.resolve_qot_column(pay)
res = Q.build_qot_rows(pay, tpl, col)

data = Q.fill_qot_template(Upload(TEMPLATE), res["rows"])
wb = openpyxl.load_workbook(io.BytesIO(data))
src = openpyxl.load_workbook(TEMPLATE)
check("sheet list preserved", wb.sheetnames, src.sheetnames)
check("Instructions kept", "Instructions" in wb.sheetnames, True)

ws = wb[Q.TEMPLATE_SHEET]
check("header row intact",
      [ws.cell(row=1, column=i + 1).value for i in range(8)], Q.TEMPLATE_HEADERS)
check("row count = emitted rows", ws.max_row - 1, 289)

out = pd.read_excel(io.BytesIO(data), sheet_name=Q.TEMPLATE_SHEET, dtype=str)
check("no blank template rows left", int(out["Pay Date"].fillna("").eq("").sum()), 0)
check("every row has all three dates",
      int(out[["Period Start Date", "Period End Date", "Pay Date"]]
          .fillna("").eq("").any(axis=1).sum()), 0)
check("total premium",
      round(pd.to_numeric(out["QOT Premium"]).sum(), 2), 13589.77)
check("first employee", out["Employee ID"].iloc[0], "11M4UJ5DM")

check("filename", Q.output_filename("CDC", Upload(TEMPLATE)),
      "CDC_Qualified_Overtime_Template_2026_filled.xlsx")

print()
print("TASK 4", "PASS" if ok else "FAIL")
sys.exit(0 if ok else 1)
