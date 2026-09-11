"""Verify ADP Time Off policy mapping (spec 2026-09-11) on real client files.

Run:  python scratch/verify_timeoff_policy_mapping.py [readers|filling|audit|ui ...]

BASELINE is origin/main before this work. With every ADP policy mapped to a
one-policy template and Salaried included, the new code must reproduce the
baseline exactly — that is the regression guarantee for one-policy clients.
"""
import importlib.util
import io
import os
import subprocess
import sys
import tempfile
from collections import defaultdict

import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from apps.adp import timeoff_audit as ta          # noqa: E402

BASELINE = "038c330"
D = r"C:\Users\rohit.kaushik\Downloads"
MOSES = (D + r"\Moses_Solutions_LLC_Time_Off_Balance_Summary.xlsx",
         D + r"\Moses_Time Off Import (2).xlsx", None)
EXPRESS = (D + r"\Express_Package_System_Inc_Time_Off_Balance_Summary.xlsx",
           D + r"\Time Off Import (1).xlsx",
           D + r"\Multi_Client_EXPRESS PACKAGE SYSTEM_Employee_Census.xlsm")
HD = (D + r"\High Distinction\New folder\Time Off Balance Summary.xlsx",
      D + r"\High Distinction\New folder\Time Off Import.xlsx",
      D + r"\High Distinction\Multi_Client_High Distinction Logistics LLC_Employee_Census.xlsm")
OLD_EXC_COLS = ["Employee ID", "Employee Name", "Issue Category", "ADP Balance"]


class U(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name
        self.size = len(data)


def up(path):
    return U(open(path, "rb").read(), os.path.basename(path))


failures = 0


def check(label, ok, detail=""):
    global failures
    print("   %-70s %s%s" % (label, "OK" if ok else "FAIL",
                             "" if ok else "  <- " + str(detail)[:300]))
    failures += 0 if ok else 1


def baseline():
    src = subprocess.run(["git", "-C", ROOT, "show", BASELINE + ":apps/adp/timeoff_audit.py"],
                         capture_output=True, text=True, encoding="utf-8", check=True).stdout
    path = os.path.join(tempfile.gettempdir(), "timeoff_audit_baseline.py")
    open(path, "w", encoding="utf-8").write(src)
    spec = importlib.util.spec_from_file_location("timeoff_audit_baseline", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def independent_balances(path):
    """(id, policy) -> balance, summed here rather than by the module under test."""
    ws = openpyxl.load_workbook(path).worksheets[0]
    head = [str(ws.cell(row=1, column=c).value or "").strip().upper()
            for c in range(1, ws.max_column + 1)]
    ci, cp, cb = (head.index(h) + 1 for h in ("ASSOCIATE ID", "POLICY NAME", "BALANCE AMOUNT"))
    out = defaultdict(float)
    for r in range(2, ws.max_row + 1):
        eid = ta.clean_id(ws.cell(row=r, column=ci).value)
        if eid:
            out[(eid, str(ws.cell(row=r, column=cp).value).strip())] += float(
                ta._evaluate_cell(ws.cell(row=r, column=cb).value))
    return {k: round(v, 2) for k, v in out.items()}


def new_run(files, mapping=None, include_salaried=False):
    """The new backend without Streamlit. Returns a dict of every intermediate."""
    a, t, _ = files
    adp_df, err = ta.read_adp_balances(up(a))
    assert not err, err
    tpl, err = ta.read_template_rows(up(t))
    assert not err, err
    auto = ta.auto_map([p for p, _, _ in ta.policy_summary(adp_df)], tpl["policies"])
    mapping = auto if mapping is None else mapping
    plan = ta.plan_fill(tpl, adp_df, mapping, include_salaried)
    wb, filled, err = ta.fill_import_template(up(t), plan["decisions"])
    assert not err, err
    return {"adp": adp_df, "tpl": tpl, "auto": auto, "mapping": mapping,
            "plan": plan, "wb": wb, "filled": filled}


def legacy_mapping(files):
    """Every ADP policy -> the template's one policy: what the old code did."""
    a, t, _ = files
    adp_df, _ = ta.read_adp_balances(up(a))
    tpl, _ = ta.read_template_rows(up(t))
    assert len(tpl["policies"]) == 1, tpl["policies"]
    return {p: tpl["policies"][0] for p, _, _ in ta.policy_summary(adp_df)}


def openings(wb_or_bytes):
    wb = (openpyxl.load_workbook(io.BytesIO(wb_or_bytes))
          if isinstance(wb_or_bytes, (bytes, bytearray)) else wb_or_bytes)
    ws = ta._template_sheet(wb)
    return {r: ws.cell(row=r, column=8).value for r in range(ta.UZIO_HEADER_ROW + 1, ws.max_row + 1)}


def sheet(xbytes, name):
    return pd.read_excel(io.BytesIO(xbytes), sheet_name=name, dtype=str).fillna("")


# ---------------------------------------------------------------- readers
def group_readers():
    print("\n== readers")
    r = new_run(MOSES)
    check("Moses ADP policies kept apart",
          set(r["adp"]["policy"]) == {"PTO", "Salary  PTO"}, set(r["adp"]["policy"]))
    summ = {p: (n, tot) for p, n, tot in ta.policy_summary(r["adp"])}
    check("Moses PTO summary is 204 EEs / 5056.63",
          summ.get("PTO") == (204, 5056.63), summ)
    check("Moses template policies in template order",
          r["tpl"]["policies"] == ["NY State Prenatal Leave", "Paid PTO"], r["tpl"]["policies"])
    check("Moses template has Pay Type; 12 Salaried rows",
          r["tpl"]["has_pay_type"] and sum(t["salaried"] for t in r["tpl"]["rows"]) == 12)
    check("auto_map Moses: both PTO policies -> Paid PTO",
          r["auto"] == {"PTO": "Paid PTO", "Salary  PTO": "Paid PTO"}, r["auto"])
    h = new_run(HD)
    check("auto_map HD: Amazon PTO -> Paid PTO, Salaried policies -> Do not import",
          h["auto"] == {"Amazon PTO": "Paid PTO",
                        "Salaried Sick Time Policy": ta.DO_NOT_IMPORT,
                        "Salaried Time Off Policy": ta.DO_NOT_IMPORT}, h["auto"])
    hansen = ta.auto_map(["New York City Prenatal Leave", "New York City Unpaid Leave", "PTO"],
                         ["NY State Prenatal Leave", "Paid PTO"])
    check("auto_map guesses PTO only, never a leave policy",
          hansen == {"New York City Prenatal Leave": ta.DO_NOT_IMPORT,
                     "New York City Unpaid Leave": ta.DO_NOT_IMPORT, "PTO": "Paid PTO"}, hansen)
    check("auto_map falls back to the single PTO-named template policy",
          ta.auto_map(["Amazon PTO"], ["PTO"]) == {"Amazon PTO": "PTO"})
    check("auto_map maps nothing when the PTO target is ambiguous",
          ta.auto_map(["PTO"], ["PTO A", "PTO B"]) == {"PTO": ta.DO_NOT_IMPORT})
    check("_sum_money never returns -0.0",
          str(ta._sum_money(pd.Series([-1e-15]))) == "0.0")


# ---------------------------------------------------------------- filling
def group_filling():
    print("\n== filling")
    base = baseline()
    r = new_run(MOSES)
    out, tpl_open = openings(r["wb"]), {t["row"]: t["opening"] for t in r["tpl"]["rows"]}
    pren = [t for t in r["tpl"]["rows"] if t["policy"] == "NY State Prenatal Leave"]
    check("Moses: all 96 prenatal rows stay 20.0",
          len(pren) == 96 and all(out[t["row"]] == 20.0 for t in pren))
    truth = independent_balances(MOSES[0])
    written = [t for t in r["tpl"]["rows"] if t["policy"] == "Paid PTO"
               and out[t["row"]] != tpl_open[t["row"]]]
    wrong = [(t["raw_id"], out[t["row"]], truth.get((t["id"], "PTO")))
             for t in r["tpl"]["rows"] if t["policy"] == "Paid PTO" and not t["salaried"]
             and not ta._is_blank(t["opening"]) and (t["id"], "PTO") in truth
             and out[t["row"]] != truth[(t["id"], "PTO")]]
    check("Moses: every assigned Hourly Paid PTO row equals ADP PTO", not wrong, wrong[:3])
    check("Moses: 86 rows written", r["filled"] == 86 and len(written) <= 86, r["filled"])
    check("Moses: Salaried rows untouched",
          all(out[t["row"]] == t["opening"] for t in r["tpl"]["rows"] if t["salaried"]))

    b_fill, _, _ = base.run_tool(up(EXPRESS[0]), up(EXPRESS[1]), up(EXPRESS[2]))
    on = new_run(EXPRESS, include_salaried=True)
    check("Express, Salaried on: filled template identical to baseline",
          openings(on["wb"]) == openings(b_fill))
    off = new_run(EXPRESS)
    base_open, off_open = openings(b_fill), openings(off["wb"])
    sal_rows = {t["row"] for t in off["tpl"]["rows"] if t["salaried"]}
    diff = {row for row in base_open if base_open[row] != off_open[row]}
    check("Express, Salaried off: every changed row is a Salaried row",
          diff <= sal_rows, sorted(diff - sal_rows)[:5])
    check("Express: 9 rows skipped as Salaried",
          sum(a == "salaried" for _, a, _ in off["plan"]["decisions"]) == 9)

    b_hd, _, _ = base.run_tool(up(HD[0]), up(HD[1]), up(HD[2]))
    leg = new_run(HD, mapping=legacy_mapping(HD), include_salaried=True)
    check("High Distinction, legacy mapping: identical to baseline",
          openings(leg["wb"]) == openings(b_hd))


# ---------------------------------------------------------------- audit
def group_audit():
    print("\n== audit")
    base = baseline()
    for name, files in (("Express", EXPRESS), ("High Distinction", HD)):
        a, t, c = files
        bf, ba, bs = base.run_tool(up(a), up(t), up(c))
        nf, na, ns = ta.run_tool(up(a), up(t), up(c), mapping=legacy_mapping(files),
                                 include_salaried=True)
        for s in ("Balance vs UZIO Status", "Unassigned Policies"):
            check("%s legacy: %s identical" % (name, s), sheet(ba, s).equals(sheet(na, s)))
        check("%s legacy: Exception Summary rows identical" % name,
              sheet(ba, "Exception Summary")[OLD_EXC_COLS].equals(
                  sheet(na, "Exception Summary")[OLD_EXC_COLS]))
        check("%s legacy: stats identical" % name,
              {k: ns[k] for k in bs} == bs, (bs, {k: ns.get(k) for k in bs}))

    _, audit, stats = ta.run_tool(up(HD[0]), up(HD[1]), up(HD[2]))
    exc = sheet(audit, "Exception Summary")
    ni = exc[exc["Issue Category"].str.startswith("ADP policy not imported")]
    got = {(r["Employee ID"], r["Issue Category"], float(r["ADP Balance"])) for _, r in ni.iterrows()}
    want = {("I40QAXG83", "ADP policy not imported (Salaried Sick Time Policy)", 4.0),
            ("I40QAXG83", "ADP policy not imported (Salaried Time Off Policy)", 1.0),
            ("2OFLJ8BJI", "ADP policy not imported (Salaried Time Off Policy)", 2.5)}
    check("HD: non-zero Do-not-import balances reported, zero ones not", got == want, got)
    bvs = sheet(audit, "Balance vs UZIO Status")
    check("HD: ADP Balance counts mapped policies only",
          not {"I40QAXG83", "2OFLJ8BJI"} & set(bvs["Employee ID"]))
    pm = sheet(audit, "Policy Mapping")
    check("HD: Policy Mapping lists all three ADP policies as Auto",
          set(pm["ADP Policy"]) == {"Amazon PTO", "Salaried Sick Time Policy",
                                    "Salaried Time Off Policy"}
          and set(pm["Set By"]) == {"Auto"}, pm.values.tolist())

    _, audit, stats = ta.run_tool(up(EXPRESS[0]), up(EXPRESS[1]), up(EXPRESS[2]))
    exc = sheet(audit, "Exception Summary")
    sal = exc[exc["Issue Category"] == "Salaried — not filled (Paid PTO)"]
    check("Express: 9 'Salaried — not filled (Paid PTO)' rows",
          len(sal) == 9 and stats["salaried"] == 9, (len(sal), stats.get("salaried")))

    r = new_run(MOSES)
    sheets, counts = ta.build_audit_sheets(up(MOSES[1]), r["tpl"], r["adp"], None,
                                           r["plan"], r["mapping"], r["auto"])
    pm = sheets["Policy Mapping"].astype(str).values.tolist()
    check("Moses: Policy Mapping rows",
          pm == [["PTO", "Paid PTO", "204", "5056.63", "Auto"],
                 ["Salary  PTO", "Paid PTO", "1", "0.0", "Auto"],
                 ["(none)", "NY State Prenatal Leave", "—", "—", "Not filled"]], pm)
    check("Moses: 112 missing, no Salaried exceptions",
          counts["missing"] == 112 and counts["salaried"] == 0, counts)
    you = dict(r["mapping"], **{"Salary  PTO": ta.DO_NOT_IMPORT})
    plan = ta.plan_fill(r["tpl"], r["adp"], you, False)
    sheets, counts = ta.build_audit_sheets(up(MOSES[1]), r["tpl"], r["adp"], None,
                                           plan, you, r["auto"])
    row = sheets["Policy Mapping"].set_index("ADP Policy").loc["Salary  PTO"]
    check("Moses: a changed mapping is 'You'; a zero balance is not reported",
          row["Set By"] == "You" and not counts["not_imported"], (row.to_dict(), counts))


# ---------------------------------------------------------------- ui
def group_ui():
    print("\n== ui")
    import json
    from streamlit.testing.v1 import AppTest

    def script():
        # AppTest re-executes this from source: no closures, and it cannot drive
        # st.file_uploader, so the uploader is stubbed by widget key.
        import io as _io
        import json as _json
        import os as _os
        import sys as _sys
        _sys.path.insert(0, _os.environ["TO_REPO_ROOT"])
        import streamlit as _st
        from apps.adp import timeoff_audit
        files = _json.loads(_os.environ.get("TO_FILES", "{}"))

        class _U(_io.BytesIO):
            def __init__(self, p):
                with open(p, "rb") as fh:
                    super().__init__(fh.read())
                self.name = _os.path.basename(p)
                self.size = len(self.getvalue())

        _st.file_uploader = lambda *a, **k: (_U(files[k["key"]]) if k.get("key") in files
                                             else None)
        timeoff_audit.render_ui()

    a, t, c = EXPRESS
    os.environ["TO_REPO_ROOT"] = ROOT
    os.environ["TO_FILES"] = json.dumps({"at_a": a, "at_u": t})
    at = AppTest.from_function(script, default_timeout=180).run()
    check("mapping renders once ADP + template are in",
          not at.exception and len(at.selectbox) == 1,
          at.exception or len(at.selectbox))
    if not len(at.selectbox):
        return
    sb = at.selectbox[0]
    check("Amazon PTO pre-mapped to Paid PTO",
          "Amazon PTO" in sb.label and sb.value == "Paid PTO", (sb.label, sb.value))
    check("Salaried checkbox starts unchecked",
          len(at.checkbox) == 1 and at.checkbox[0].value is False)

    os.environ["TO_FILES"] = json.dumps({"at_a": a, "at_u": t, "at_c": c})
    at.run()
    at.button(key="run_timeoff_adp").click().run()
    metrics = {m.label: m.value for m in at.metric}
    expect = new_run(EXPRESS)["filled"]
    check("Generate writes the expected balances",
          len(at.get("download_button")) == 2
          and str(metrics.get("Balances written")) == str(expect),
          (len(at.get("download_button")), metrics, expect))
    check("left-out warning mentions the Salaried rows",
          any("Salaried" in w.value for w in at.warning), [w.value for w in at.warning])

    at.selectbox[0].select(ta.DO_NOT_IMPORT).run()
    check("changing a mapping discards the previous downloads",
          len(at.get("download_button")) == 0)
    at.button(key="run_timeoff_adp").click().run()
    metrics = {m.label: m.value for m in at.metric}
    check("nothing mapped -> nothing written",
          str(metrics.get("Balances written")) == "0", metrics)


GROUPS = {"readers": group_readers, "filling": group_filling, "audit": group_audit,
          "ui": group_ui}


def main(argv):
    for name in (argv or list(GROUPS)):
        try:
            GROUPS[name]()
        except Exception as e:                       # a missing function is a FAIL
            check("group %s ran" % name, False, repr(e))
    print()
    print("FAIL - %d check(s)" % failures if failures else "PASS - every check held")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
