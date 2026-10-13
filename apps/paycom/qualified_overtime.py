import io
import os
import re
from decimal import Decimal, InvalidOperation

import openpyxl
import pandas as pd
import streamlit as st

from utils.file_io import read_report

APP_TITLE = "Paycom - Qualified Overtime Wages"

# =========================================================
# Paycom - Qualified Overtime Wages
# - Input: one or more Paycom "Estimated Qualified Premiums" reports, and the
#   UZIO Qualified Overtime template.
# - Paycom's report holds ONE row per employee with the year's premium and no
#   pay periods, so every row written is DATELESS. Template rule 9: a dateless
#   row is the employee's cumulative total for the tax year.
# - A premium of 0 is skipped, never written: rule 9 also says a dateless 0
#   clears everything already imported for that employee.
# - Rule 11: an employee left out of the file is not changed, so writing only
#   the employees Paycom reports is safe.
# =========================================================

# Paycom report columns (matched case-insensitively, spaces trimmed).
PAYCOM_ID_COL = "EE Code"
PAYCOM_PREMIUM_COL = "Estimated Qualified Premiums"
PAYCOM_NAME_COL = "Employee Name"
PAYCOM_STATUS_COL = "Status"
PAYCOM_DESC_COL = "W-2 Box Description"

# UZIO Qualified Overtime template. Headers sit on row 1 of "QOT Details";
# "Instructions" is a second sheet that must survive into the output untouched.
TEMPLATE_SHEET = "QOT Details"
TEMPLATE_HEADERS = ["Employee ID", "First Name", "Last Name", "Employment Status",
                    "Period Start Date", "Period End Date", "Pay Date", "QOT Premium"]


def _find_col(columns, name):
    want = name.strip().lower()
    return next((c for c in columns if str(c).strip().lower() == want), None)


def parse_premium(v):
    """A premium as Decimal (2 places). Blank -> None. Not a number -> raises ValueError.

    "$1,234.50" -> 1234.50; "(12.34)" -> -12.34 (accounting negative).
    """
    s = "" if v is None else str(v).strip()
    if s == "" or s.lower() in ("nan", "none"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    t = s.strip("()").replace("$", "").replace(",", "").strip()
    try:
        d = Decimal(t)
    except InvalidOperation:
        raise ValueError(s)
    d = -d if neg else d
    return d.quantize(Decimal("0.01"))


# ---------------------------------------------------------------- Paycom side

def _read_one(file, label):
    """One Paycom report as text. Returns (df, error).

    `keep_default_na=False`: pandas would otherwise turn "N/A" or "NA" in the
    premium column into a blank, and a blank premium is silently skipped — a
    junk value must reach parse_premium and block the run instead.
    """
    name = label
    try:
        df = read_report(file, header=0, dtype=str, keep_default_na=False)
    except Exception as e:
        return None, "%s — could not be read: %s" % (name, e)
    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in (PAYCOM_ID_COL, PAYCOM_PREMIUM_COL) if not _find_col(df.columns, c)]
    if missing:
        return None, "%s — missing column(s): %s" % (name, ", ".join(missing))
    df = df.rename(columns={_find_col(df.columns, c): c
                            for c in (PAYCOM_ID_COL, PAYCOM_PREMIUM_COL, PAYCOM_NAME_COL,
                                      PAYCOM_STATUS_COL, PAYCOM_DESC_COL)
                            if _find_col(df.columns, c)})
    df = df.fillna("")
    df[PAYCOM_ID_COL] = df[PAYCOM_ID_COL].astype(str).str.strip()
    df = df[df[PAYCOM_ID_COL] != ""].copy()
    df["_file"] = name
    return df, None


def read_paycom_files(files):
    """Every uploaded report, stacked. Returns (DataFrame, errors).

    A file that cannot be read, or lacks EE Code / Estimated Qualified Premiums,
    is reported and left out; the others still count.

    Each upload is told apart by its position, not its name: the same report
    uploaded twice has the same name, and treating the two as one file would
    ADD their premiums instead of blocking the double count.
    """
    files = list(files or [])
    names = [getattr(f, "name", "") or "upload" for f in files]
    frames, errors = [], []
    for i, (f, name) in enumerate(zip(files, names), start=1):
        label = name if names.count(name) == 1 else "%s (upload %d)" % (name, i)
        df, err = _read_one(f, label)
        if err:
            errors.append(err)
        else:
            frames.append(df)
    return (pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()), errors


# ---------------------------------------------------------------- template

def read_qot_template(file):
    """The template's QOT Details sheet as strings. Returns (df, error)."""
    file.seek(0)
    try:
        book = pd.read_excel(io.BytesIO(file.read()), sheet_name=None, dtype=str)
    except Exception as e:
        return None, "Could not read the template: %s" % e
    if TEMPLATE_SHEET not in book:
        return None, ("The template has no `%s` sheet. Sheets found: %s."
                      % (TEMPLATE_SHEET, ", ".join(book) or "none"))
    df = book[TEMPLATE_SHEET].copy()
    df.columns = [str(c).strip() for c in df.columns]
    missing = [h for h in TEMPLATE_HEADERS if h not in df.columns]
    if missing:
        return None, ("`%s` is missing column(s): %s."
                      % (TEMPLATE_SHEET, ", ".join(missing)))
    df = df[TEMPLATE_HEADERS].copy()
    df["Employee ID"] = df["Employee ID"].fillna("").astype(str).str.strip()
    df = df[df["Employee ID"] != ""].reset_index(drop=True)
    return df.fillna(""), None


# ---------------------------------------------------------------- rows

def build_qot_rows(paycom_df, template_df):
    """One dateless row per employee with a non-zero premium.

    - Several rows for one employee in ONE file are added together: the
      template allows a single dateless (cumulative) row per employee.
    - The same employee in TWO files blocks the run — adding them could count
      the same overtime twice, and which file is right is not ours to guess.
    - A premium that is not a number blocks the run rather than vanishing.
    - 0 (after adding) is skipped: a dateless 0 would clear the employee's
      imported QOT (rule 9).
    - An employee missing from the template is left out and reported.
    """
    out = {"rows": [], "bad_values": [], "cross_file": [], "combined": [],
           "not_in_template": [], "status_mismatch": [], "zero_skipped": [],
           "total_qot": Decimal("0.00"), "descriptions": {}}
    if paycom_df.empty:
        return out

    df = paycom_df.copy()
    amounts = []
    for _, r in df.iterrows():
        try:
            amounts.append(parse_premium(r[PAYCOM_PREMIUM_COL]))
        except ValueError:
            amounts.append(None)
            out["bad_values"].append({"File": r["_file"], "EE Code": r[PAYCOM_ID_COL],
                                      "Value": r[PAYCOM_PREMIUM_COL]})
    df["_amt"] = amounts
    if PAYCOM_DESC_COL in df.columns:
        out["descriptions"] = df[PAYCOM_DESC_COL].str.strip().replace("", "(blank)") \
            .value_counts().to_dict()

    files_per_emp = df.groupby(PAYCOM_ID_COL)["_file"].agg(lambda s: sorted(set(s)))
    for eid, files in files_per_emp.items():
        if len(files) > 1:
            out["cross_file"].append({"EE Code": eid, "Files": ", ".join(files)})

    names = template_df.set_index("Employee ID")[
        ["First Name", "Last Name", "Employment Status"]].to_dict("index")

    for eid, g in df.groupby(PAYCOM_ID_COL, sort=True):
        vals = [a for a in g["_amt"] if a is not None]
        total = sum(vals, Decimal("0.00"))
        if len(g) > 1 and len(set(g["_file"])) == 1:
            out["combined"].append({
                "EE Code": eid, "Rows": len(g),
                "Descriptions": ", ".join(g.get(PAYCOM_DESC_COL, pd.Series([""] * len(g)))
                                          .astype(str).str.strip()),
                "Combined Premium": float(total)})
        if total == 0:
            out["zero_skipped"].append(eid)
            continue
        if eid not in names:
            out["not_in_template"].append({"EE Code": eid,
                                           "Employee Name": g.iloc[0].get(PAYCOM_NAME_COL, ""),
                                           "Premium": float(total)})
            continue
        n = names[eid]
        pstatus = str(g.iloc[0].get(PAYCOM_STATUS_COL, "")).strip()
        if pstatus and pstatus.lower() != str(n["Employment Status"]).strip().lower():
            out["status_mismatch"].append({"EE Code": eid, "Paycom Status": pstatus,
                                           "Template Status": n["Employment Status"]})
        out["rows"].append({
            "Employee ID": eid,
            "First Name": n["First Name"],
            "Last Name": n["Last Name"],
            "Employment Status": n["Employment Status"],
            "Period Start Date": None,
            "Period End Date": None,
            "Pay Date": None,
            "QOT Premium": float(total),
        })
        out["total_qot"] += total
    out["total_qot"] = float(out["total_qot"])
    return out


def is_blocked(result):
    """A non-numeric premium or an employee in two files makes the output unsafe."""
    return bool(result["bad_values"] or result["cross_file"])


# ---------------------------------------------------------------- output

def fill_qot_template(template_file, rows):
    """Write the rows into a copy of the uploaded template. Returns .xlsx bytes.

    openpyxl keeps the Instructions sheet and the formatting. The template's
    pre-filled rows (every employee, blank dates, blank premium) are cleared
    first, so only the employees Paycom reports are imported; rule 11 leaves
    everyone else as they are.
    """
    template_file.seek(0)
    wb = openpyxl.load_workbook(io.BytesIO(template_file.read()))
    ws = wb[TEMPLATE_SHEET]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for i, r in enumerate(rows, start=2):
        for j, h in enumerate(TEMPLATE_HEADERS, start=1):
            ws.cell(row=i, column=j).value = r[h]
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def output_filename(client_name, template_file):
    """<Client>_<uploaded template's name>_filled.xlsx"""
    base = os.path.splitext(getattr(template_file, "name", "") or "template")[0]
    client = re.sub(r"\s+", " ", str(client_name or "Client")).strip() or "Client"
    return "%s_%s_filled.xlsx" % (client, base)


# ---------------------------------------------------------------- UI

def render_ui():
    st.title(APP_TITLE)
    st.markdown("""
    **Upload**
    1. **Paycom Estimated Qualified Premiums report(s)** — one or more (.xlsx, .xls, .csv)
    2. **UZIO Qualified Overtime template** (.xlsx)

    Paycom reports one premium per employee for the year, with no pay periods, so
    every row is written **without dates** — UZIO reads a dateless row as the
    employee's total for the tax year. A premium of **0 is not written**: a dateless 0
    would clear what is already imported for that employee. Employees left out of the
    file are not changed in UZIO.
    """)

    client_name = st.text_input("Client Name", value="Client", key="pqot_client")
    c1, c2 = st.columns(2)
    with c1:
        pay_files = st.file_uploader("Paycom Estimated Qualified Premiums report(s)",
                                     type=["xlsx", "xls", "csv"],
                                     accept_multiple_files=True, key="pqot_pay")
    with c2:
        tpl_file = st.file_uploader("UZIO Qualified Overtime template",
                                    type=["xlsx"], key="pqot_tpl")

    if not pay_files or not tpl_file:
        st.info("Upload the Paycom report(s) and the UZIO template to begin.")
        return

    paycom, read_errors = read_paycom_files(pay_files)
    for e in read_errors:
        st.warning(e)
    if paycom.empty:
        st.error("None of the uploaded Paycom reports could be used.")
        return

    template, tpl_err = read_qot_template(tpl_file)
    if tpl_err:
        st.error(tpl_err)
        return

    result = build_qot_rows(paycom, template)

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Rows to import", len(result["rows"]))
    m2.metric("Employees in Paycom", paycom[PAYCOM_ID_COL].nunique())
    m3.metric("Total QOT", "%.2f" % result["total_qot"])
    m4.metric("Skipped (premium 0)", len(result["zero_skipped"]))

    if result["bad_values"]:
        st.error("**%d premium(s) are not a number**, so no file was produced. Fix them "
                 "in the Paycom report and upload it again." % len(result["bad_values"]))
        st.dataframe(pd.DataFrame(result["bad_values"]), hide_index=True, width="stretch")

    if result["cross_file"]:
        st.error("**%d employee(s) appear in more than one uploaded file**, so no file was "
                 "produced. Adding them together could count the same overtime twice — "
                 "upload only the report that is right for each employee."
                 % len(result["cross_file"]))
        st.dataframe(pd.DataFrame(result["cross_file"]), hide_index=True, width="stretch")

    if result["combined"]:
        st.info("**%d employee(s) have more than one row in a report** — their premiums "
                "were added into one row, because UZIO takes a single total per employee."
                % len(result["combined"]))
        st.dataframe(pd.DataFrame(result["combined"]), hide_index=True, width="stretch")

    if result["not_in_template"]:
        st.warning("**%d employee(s) have a premium but no row in the template.** They were "
                   "left out — UZIO has no such employee to import against."
                   % len(result["not_in_template"]))
        st.dataframe(pd.DataFrame(result["not_in_template"]), hide_index=True, width="stretch")

    if result["status_mismatch"]:
        st.warning("**%d employee(s) have a different status in Paycom and in the template.** "
                   "Nothing was changed — UZIO reads only the Employee ID on import — but "
                   "worth a look." % len(result["status_mismatch"]))
        st.dataframe(pd.DataFrame(result["status_mismatch"]), hide_index=True, width="stretch")

    if result["zero_skipped"]:
        st.caption("Premium 0, not written: " + ", ".join(result["zero_skipped"]))

    if is_blocked(result):
        return
    if not result["rows"]:
        st.info("No qualified overtime found in the uploaded report(s).")
        return

    st.markdown("### Rows to be written")
    st.dataframe(pd.DataFrame(result["rows"], columns=TEMPLATE_HEADERS),
                 hide_index=True, width="stretch")

    # st.download_button reruns the script, so the bytes are built here (outside
    # any button block) and the widget simply serves them.
    st.download_button(
        "⬇️ Download filled template",
        data=fill_qot_template(tpl_file, result["rows"]),
        file_name=output_filename(client_name, tpl_file),
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        key="pqot_dl",
    )
