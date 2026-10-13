"""Verify the Paycom Qualified Overtime Wages tool on the real Banda files.

Run:  python scratch/verify_paycom_qot.py [banda|synthetic|ui ...]

Spec: docs/superpowers/specs/2026-10-13-paycom-qot-design.md
"""
import io
import json
import os
import sys

import openpyxl
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from apps.paycom import qualified_overtime as q          # noqa: E402

B = r"C:\Users\rohit.kaushik\Downloads\Banda Logistics"
REPORT = B + r"\20261013014052_estimated_qualified_premiums_report_ad2451e0.xlsx"
TEMPLATE = B + r"\Qualified_Overtime_Template_2026.xlsx"

failures = 0


class U(io.BytesIO):
    def __init__(self, data, name):
        super().__init__(data)
        self.name = name
        self.size = len(data)


def up(path_or_bytes, name=None):
    if isinstance(path_or_bytes, bytes):
        return U(path_or_bytes, name or "upload.xlsx")
    return U(open(path_or_bytes, "rb").read(), os.path.basename(path_or_bytes))


def check(label, ok, detail=""):
    global failures
    print("   %-72s %s%s" % (label, "OK" if ok else "FAIL",
                             "" if ok else "  <- " + str(detail)[:300]))
    failures += 0 if ok else 1


def run(files, template=TEMPLATE):
    paycom, errors = q.read_paycom_files(files)
    tpl, err = q.read_qot_template(up(template))
    assert not err, err
    return paycom, errors, tpl, q.build_qot_rows(paycom, tpl)


def report_bytes(rows, header=None):
    header = header or ["Client Code", "EE Code", "Employee Name", "Status", "W-2 Box",
                        "W-2 Box Description", "Estimated Qualified Premiums", "Source"]
    buf = io.BytesIO()
    pd.DataFrame(rows, columns=header).to_excel(buf, index=False)
    return buf.getvalue()


# ---------------------------------------------------------------- banda
def group_banda():
    print("\n== Banda: real report + real 2026 template")
    raw = pd.read_excel(REPORT, dtype=str)
    raw_amt = pd.to_numeric(raw["Estimated Qualified Premiums"])
    paycom, errors, tpl, res = run([up(REPORT)])
    check("report read with no errors", not errors, errors)
    nonzero = raw[raw_amt != 0]
    check("rows = employees with a non-zero premium (%d)" % len(nonzero),
          len(res["rows"]) == len(nonzero), len(res["rows"]))
    check("total QOT = report total (%.2f)" % raw_amt.sum(),
          round(res["total_qot"], 2) == round(raw_amt.sum(), 2), res["total_qot"])
    check("the 2 zero-premium employees are skipped",
          sorted(res["zero_skipped"]) == sorted(raw[raw_amt == 0]["EE Code"]), res["zero_skipped"])
    by_id = dict(zip(raw["EE Code"], raw_amt))
    check("every row's premium equals Paycom's",
          all(abs(r["QOT Premium"] - by_id[r["Employee ID"]]) < 1e-9 for r in res["rows"]))
    check("every row is dateless",
          all(r["Period Start Date"] is None and r["Period End Date"] is None
              and r["Pay Date"] is None for r in res["rows"]))
    check("names and status come from the template",
          all(r["First Name"] for r in res["rows"]))
    check("2 status mismatches reported (A0DX, A0JK)",
          sorted(m["EE Code"] for m in res["status_mismatch"]) == ["A0DX", "A0JK"],
          res["status_mismatch"])
    check("nothing blocked, nobody missing from the template",
          not q.is_blocked(res) and not res["not_in_template"])

    out = openpyxl.load_workbook(io.BytesIO(q.fill_qot_template(up(TEMPLATE), res["rows"])))
    src = openpyxl.load_workbook(TEMPLATE)
    check("sheets kept: %s" % src.sheetnames, out.sheetnames == src.sheetnames)
    same_instr = all(a.value == b.value for ra, rb in zip(src["Instructions"].iter_rows(),
                                                          out["Instructions"].iter_rows())
                     for a, b in zip(ra, rb))
    check("Instructions sheet unchanged", same_instr)
    ws = out[q.TEMPLATE_SHEET]
    body = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
    check("QOT Details holds only the %d written rows (330 pre-filled cleared)" % len(res["rows"]),
          len(body) == len(res["rows"]), len(body))
    check("header row unchanged",
          [c.value for c in ws[1]][:8] == q.TEMPLATE_HEADERS)
    check("dates empty and premium numeric in the file",
          all(r[4] is None and r[5] is None and r[6] is None and isinstance(r[7], (int, float))
              for r in body))
    reread = pd.read_excel(io.BytesIO(q.fill_qot_template(up(TEMPLATE), res["rows"])),
                           sheet_name=q.TEMPLATE_SHEET)
    check("file total = %.2f" % raw_amt.sum(), round(reread["QOT Premium"].sum(), 2)
          == round(raw_amt.sum(), 2), reread["QOT Premium"].sum())
    check("output name", q.output_filename("Banda", up(TEMPLATE))
          == "Banda_Qualified_Overtime_Template_2026_filled.xlsx")


# ---------------------------------------------------------------- synthetic
def group_synthetic():
    print("\n== parse_premium")
    for v, want in (("3.06", "3.06"), ("$1,234.5", "1234.50"), ("(12.34)", "-12.34"),
                    ("-5", "-5.00"), ("", None), (" ", None), (None, None)):
        got = q.parse_premium(v)
        check("%-10r -> %s" % (v, want), (str(got) if got is not None else None) == want, got)
    try:
        q.parse_premium("abc")
        check("'abc' raises", False)
    except ValueError:
        check("'abc' raises", True)

    print("\n== two rows for one employee in one file are added")
    f1 = report_bytes([["X", "A009", "BLUNT, JAIDAH", "ACTIVE", "[14] Other", "Qualified OT", "3.06", "T"],
                       ["X", "A009", "BLUNT, JAIDAH", "ACTIVE", "[14] Other", "Qualified Tips", "10.00", "T"],
                       ["X", "A00N", "MAINARDI, DAVID", "ACTIVE", "[14] Other", "Qualified OT", "(2.00)", "T"]])
    _, _, _, res = run([up(f1, "one.xlsx")])
    rows = {r["Employee ID"]: r["QOT Premium"] for r in res["rows"]}
    check("A009 = 3.06 + 10.00 = 13.06, one row", rows.get("A009") == 13.06 and
          sum(1 for r in res["rows"] if r["Employee ID"] == "A009") == 1, rows)
    check("combine reported with both descriptions",
          res["combined"] and "Qualified Tips" in res["combined"][0]["Descriptions"], res["combined"])
    check("accounting negative kept: A00N = -2.00", rows.get("A00N") == -2.0, rows)
    check("not blocked", not q.is_blocked(res))

    print("\n== same employee in two files blocks")
    f2 = report_bytes([["Y", "A009", "BLUNT, JAIDAH", "ACTIVE", "[14] Other", "Qualified OT", "1.00", "T"]])
    _, _, _, res = run([up(f1, "one.xlsx"), up(f2, "two.xlsx")])
    check("blocked, naming both files",
          q.is_blocked(res) and res["cross_file"][0]["Files"] == "one.xlsx, two.xlsx", res["cross_file"])

    print("\n== a junk premium blocks; a missing employee is reported")
    f3 = report_bytes([["X", "A009", "B, J", "ACTIVE", "", "Qualified OT", "N/A", "T"],
                       ["X", "ZZ99", "NOBODY, X", "ACTIVE", "", "Qualified OT", "5.00", "T"]])
    _, _, _, res = run([up(f3, "three.xlsx")])
    check("'N/A' blocks the run", q.is_blocked(res) and res["bad_values"][0]["Value"] == "N/A",
          res["bad_values"])
    check("ZZ99 reported as not in template",
          [r["EE Code"] for r in res["not_in_template"]] == ["ZZ99"])

    print("\n== a report without the premium column is refused")
    f4 = report_bytes([["X", "A009", "B, J"]], header=["Client Code", "EE Code", "Employee Name"])
    paycom, errors, _, _ = run([up(f4, "bad.xlsx")])
    check("error names the file and the column",
          paycom.empty and errors and "bad.xlsx" in errors[0]
          and "Estimated Qualified Premiums" in errors[0], errors)

    print("\n== the same report as CSV")
    csv = pd.read_excel(REPORT, dtype=str).to_csv(index=False).encode("utf-8")
    _, _, _, a = run([up(csv, "r.csv")])
    _, _, _, b = run([up(REPORT)])
    check("CSV gives the same rows", a["rows"] == b["rows"])


# ---------------------------------------------------------------- ui
def group_ui():
    print("\n== the screen")
    from streamlit.testing.v1 import AppTest

    def script():
        import io as _io
        import json as _json
        import os as _os
        import sys as _sys
        _sys.path.insert(0, _os.environ["PQ_ROOT"])
        import streamlit as _st
        from apps.paycom import qualified_overtime as m
        files = _json.loads(_os.environ["PQ_FILES"])

        class _U(_io.BytesIO):
            def __init__(self, p):
                with open(p, "rb") as fh:
                    super().__init__(fh.read())
                self.name = _os.path.basename(p)
                self.size = len(self.getvalue())

        def fake(*a, **k):
            v = files.get(k.get("key"))
            if v is None:
                return None
            return [_U(p) for p in v] if isinstance(v, list) else _U(v)

        _st.file_uploader = fake
        m.render_ui()

    os.environ["PQ_ROOT"] = ROOT
    os.environ["PQ_FILES"] = json.dumps({"pqot_pay": [REPORT], "pqot_tpl": TEMPLATE})
    at = AppTest.from_function(script, default_timeout=180).run()
    check("no exception", not at.exception, [e.value for e in at.exception])
    vals = {mt.label: mt.value for mt in at.metric}
    check("metrics: 94 rows, 96 employees, 6402.50, 2 skipped",
          vals.get("Rows to import") == "94" and vals.get("Employees in Paycom") == "96"
          and vals.get("Total QOT") == "6402.50" and vals.get("Skipped (premium 0)") == "2", vals)
    check("status-mismatch warning shown", any("different status" in w.value for w in at.warning))
    check("one download button", len(at.get("download_button")) == 1)

    os.environ["PQ_FILES"] = json.dumps({"pqot_pay": [REPORT, REPORT], "pqot_tpl": TEMPLATE})
    at = AppTest.from_function(script, default_timeout=180).run()
    check("the same report uploaded twice is blocked (no download)",
          any("more than one uploaded file" in e.value for e in at.error)
          and not at.get("download_button"), [e.value for e in at.error])


GROUPS = {"banda": group_banda, "synthetic": group_synthetic, "ui": group_ui}


def main(argv):
    for name in (argv or list(GROUPS)):
        try:
            GROUPS[name]()
        except Exception as e:
            import traceback
            traceback.print_exc()
            check("group %s ran" % name, False, repr(e))
    print()
    print("FAIL - %d check(s)" % failures if failures else "PASS - every check held")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
